from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from openpyxl import load_workbook
import time

# Load Excel data
workbook = load_workbook("cogmento_login_data.xlsx")
sheet = workbook.active

# Initialize Chrome
driver = webdriver.Chrome()
driver.maximize_window()
wait = WebDriverWait(driver, 5)

print("DATA-DRIVEN LOGIN TEST STARTED")

for row in range(2, sheet.max_row + 1):
    test_id = sheet.cell(row, 1).value
    email = sheet.cell(row, 2).value
    password = sheet.cell(row, 3).value
    expected = sheet.cell(row, 4).value

    driver.get("https://ui.cogmento.com/")
    time.sleep(2)  # Let page load

    try:
        # Wait for email field
        email_field = wait.until(EC.visibility_of_element_located((By.NAME, "email")))
        email_field.clear()
        email_field.send_keys(email)

        # Wait for password field
        password_field = wait.until(EC.visibility_of_element_located((By.NAME, "password")))
        password_field.clear()
        password_field.send_keys(password)

        # Click login button
        login_button = wait.until(EC.element_to_be_clickable((By.XPATH, "//div[text()='Login']")))
        login_button.click()

        # --- Check success or failure ---
        if expected == "SUCCESS":
            # Wait for the user label with "AITU" to appear
            user_label = wait.until(EC.visibility_of_element_located(
                (By.XPATH, "//b[text()='AITU']")))
            assert user_label.is_displayed()
            result = "PASSED"
            print(f"{test_id}: {result}")

            # Logout to reset
            user_menu = wait.until(EC.element_to_be_clickable((By.CSS_SELECTOR, "div.ui.dropdown")))
            user_menu.click()
            logout_button = wait.until(EC.element_to_be_clickable((By.XPATH, "//span[text()='Log Out']")))
            logout_button.click()
            wait.until(EC.visibility_of_element_located((By.NAME, "email")))

        else:
            # Wait for error message
            error_msg = wait.until(EC.visibility_of_element_located((By.CLASS_NAME, "ui.error.message")))
            assert error_msg.is_displayed()
            result = "PASSED"
            print(f"{test_id}: {result}")

    except Exception as e:
        result = "FAILED"
        print(f"{test_id} Exception: {str(e)}")
        print(f"{test_id}: {result}")

driver.quit()
print("TEST EXECUTION FINISHED")
